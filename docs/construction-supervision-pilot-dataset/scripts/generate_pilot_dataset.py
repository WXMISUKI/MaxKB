# coding=utf-8
"""
Generate a simulated construction-supervision review dataset for MaxKB pilot validation.

The generated files are intentionally fictional and safe to use for local knowledge-base tests.
"""

from __future__ import annotations

import hashlib
import json
import shutil
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[2] / "simulated-pilot-dataset" / "NJDL-JD-A1"

BASE_META = {
    "organization_id": "org-supervision-demo",
    "project_id": "project-njdl-jd-a1",
    "project_name": "南江至东岭高速公路改扩建工程",
    "contract_package_id": "contract-jd-a1",
    "contract_package_name": "JD-A1 土建施工合同段",
    "section_id": "section-k12-k18",
    "supervision_section_id": "supervision-jl-01",
    "supervision_section_name": "JL-01 监理合同段",
    "subcontract_team_id": "team-lj-01",
    "subcontract_team_name": "LJ-01 路基土石方分包作业队",
    "basis_version_id": "basis-njdl-2026-01",
    "effective_status": "sample",
    "effective_date": "2026-07-18",
}


@dataclass(frozen=True)
class DocSpec:
    doc_id: str
    folder: str
    filename: str
    title: str
    document_type: str
    review_task_id: str
    sections: list[tuple[str, list[str]]]
    tables: list[tuple[str, list[str], list[list[str]]]] | None = None


@dataclass(frozen=True)
class SheetSpec:
    doc_id: str
    folder: str
    filename: str
    title: str
    document_type: str
    review_task_id: str
    sheets: list[tuple[str, list[str], list[list[object]]]]


def ensure_dirs() -> None:
    if ROOT.exists():
        shutil.rmtree(ROOT)
    for folder in [
        "00_manifest",
        "01_project_basis",
        "02_contract_and_supervision",
        "03_team_lj01_opening_condition",
        "04_team_lj01_construction_plan",
        "05_review_examples",
        "06_source_norms",
    ]:
        (ROOT / folder).mkdir(parents=True, exist_ok=True)


def set_doc_styles(doc: Document) -> None:
    styles = doc.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"].font.size = Pt(10.5)
    styles["Heading 1"].font.name = "Microsoft YaHei"
    styles["Heading 1"].font.size = Pt(18)
    styles["Heading 1"].font.color.rgb = RGBColor(31, 78, 121)
    styles["Heading 2"].font.name = "Microsoft YaHei"
    styles["Heading 2"].font.size = Pt(14)
    styles["Heading 2"].font.color.rgb = RGBColor(31, 78, 121)


def add_metadata_table(doc: Document, meta: dict[str, str]) -> None:
    doc.add_heading("资料元数据", level=1)
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "字段"
    table.rows[0].cells[1].text = "值"
    for key, value in meta.items():
        row = table.add_row().cells
        row[0].text = key
        row[1].text = str(value)


def add_footer(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("模拟资料，仅用于 MaxKB 本地知识库验证")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(127, 127, 127)


def write_docx(spec: DocSpec) -> dict[str, str]:
    target = ROOT / spec.folder / spec.filename
    meta = {
        **BASE_META,
        "document_id": spec.doc_id,
        "document_type": spec.document_type,
        "review_task_id": spec.review_task_id,
        "source_object_id": f"source-{spec.doc_id.lower()}",
    }
    doc = Document()
    set_doc_styles(doc)
    title = doc.add_heading(spec.title, level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("本文件为虚构模拟资料，用于验证项目级知识库、metadata 过滤、开工条件审查和施工方案审查能力。")
    add_metadata_table(doc, meta)
    for heading, paragraphs in spec.sections:
        doc.add_heading(heading, level=1)
        for paragraph in paragraphs:
            doc.add_paragraph(paragraph)
    if spec.tables:
        for table_title, headers, rows in spec.tables:
            doc.add_heading(table_title, level=2)
            table = doc.add_table(rows=1, cols=len(headers))
            table.style = "Table Grid"
            for idx, header in enumerate(headers):
                table.rows[0].cells[idx].text = header
            for row_values in rows:
                row = table.add_row().cells
                for idx, value in enumerate(row_values):
                    row[idx].text = value
    add_footer(doc)
    doc.save(target)
    return manifest_entry(target, spec.doc_id, spec.document_type, spec.review_task_id)


def autosize(ws) -> None:
    for col_idx, column_cells in enumerate(ws.columns, 1):
        max_len = 10
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, min(len(value) + 2, 36))
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        ws.column_dimensions[get_column_letter(col_idx)].width = max_len


def write_xlsx(spec: SheetSpec) -> dict[str, str]:
    target = ROOT / spec.folder / spec.filename
    meta = {
        **BASE_META,
        "document_id": spec.doc_id,
        "document_type": spec.document_type,
        "review_task_id": spec.review_task_id,
        "source_object_id": f"source-{spec.doc_id.lower()}",
    }
    wb = Workbook()
    ws = wb.active
    ws.title = "metadata"
    ws.append(["字段", "值"])
    for key, value in meta.items():
        ws.append([key, value])
    for row in ws.iter_rows(min_row=1, max_row=1):
        for cell in row:
            cell.fill = PatternFill("solid", fgColor="1F4E79")
    autosize(ws)

    for sheet_name, headers, rows in spec.sheets:
        ws = wb.create_sheet(sheet_name)
        ws.append(headers)
        for row in rows:
            ws.append(row)
        ws.freeze_panes = "A2"
        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor="1F4E79")
        autosize(ws)
    wb.save(target)
    return manifest_entry(target, spec.doc_id, spec.document_type, spec.review_task_id)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as file:
        for chunk in iter(lambda: file.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def manifest_entry(path: Path, doc_id: str, document_type: str, review_task_id: str) -> dict[str, str]:
    return {
        **BASE_META,
        "document_id": doc_id,
        "document_type": document_type,
        "review_task_id": review_task_id,
        "relative_path": str(path.relative_to(ROOT)).replace("\\", "/"),
        "file_format": path.suffix.lstrip("."),
        "content_hash": sha256(path),
        "indexed_at": "",
    }


DOCS = [
    DocSpec(
        "DOC-001",
        "01_project_basis",
        "01_项目概况与合同段划分说明.docx",
        "项目概况与合同段划分说明",
        "project_basis",
        "project-context",
        [
            ("项目概况", ["南江至东岭高速公路改扩建工程为既有双向四车道高速改扩建项目，试点范围为 JD-A1 土建施工合同段。", "本模拟资料包聚焦 K12+000-K18+500 路基清表、挖方、填筑及临时排水工程。"]),
            ("合同段划分", ["JD-A1 合同段由施工总承包单位负责组织实施，JL-01 监理合同段负责施工准备、开工条件、施工方案、质量安全和资料归档审查。"]),
            ("知识库使用边界", ["本项目知识库用于辅助召回依据、资料和风险点，不直接形成正式监理审批结论。"]),
        ],
    ),
    DocSpec(
        "DOC-002",
        "02_contract_and_supervision",
        "02_JD-A1合同段主要合同条款摘编.docx",
        "JD-A1 合同段主要合同条款摘编",
        "contract_basis",
        "project-context",
        [
            ("开工条件条款", ["分包队伍进场前应提交资质、人员、设备、施工组织、安全教育、技术交底和现场准备资料，经监理审核后方可开工。"]),
            ("施工方案条款", ["涉及路基填筑、特殊路段处理、雨季施工和临时交通组织的专项方案，应在实施前完成施工单位技术负责人审批和总监理工程师审查。"]),
            ("资料归档条款", ["所有报审表、附件、审查意见、整改回复和复核记录应按项目资料命名规则归档。"]),
        ],
    ),
    DocSpec(
        "DOC-003",
        "02_contract_and_supervision",
        "03_JL-01监理规划摘编.docx",
        "JL-01 监理规划摘编",
        "supervision_basis",
        "project-context",
        [
            ("监理工作范围", ["JL-01 监理合同段负责 JD-A1 合同段施工准备、施工过程、质量安全、计量资料和竣工资料监理工作。"]),
            ("开工审查职责", ["专业监理工程师对人员、机械、材料、方案、现场条件进行初审，总监理工程师对开工条件作最终签认。"]),
            ("方案审查职责", ["施工方案审查应重点核查适用范围、依据版本、工艺流程、质量控制、安全措施和应急预案。"]),
        ],
    ),
    DocSpec(
        "DOC-004",
        "02_contract_and_supervision",
        "04_开工条件审查实施细则.docx",
        "开工条件审查实施细则",
        "review_rule",
        "opening-condition-lj01",
        [
            ("审查目标", ["确认 LJ-01 路基土石方分包作业队具备进场施工的人员、设备、材料、方案、安全和现场条件。"]),
            ("必审材料", ["企业资质、安全生产许可证、人员任命和证书、特种作业人员清单、机械设备清单、仪器校验证明、安全教育和技术交底记录。"]),
            ("审查结论", ["资料齐备且现场条件满足要求时可建议通过；存在关键岗位缺失、设备未进场、方案未批复或安全教育缺失时应退回整改。"]),
        ],
    ),
    DocSpec(
        "DOC-005",
        "02_contract_and_supervision",
        "05_施工方案审查实施细则.docx",
        "施工方案审查实施细则",
        "review_rule",
        "construction-plan-lj01",
        [
            ("审查范围", ["本细则适用于 K12+000-K18+500 路基填筑施工方案审查。"]),
            ("审查重点", ["重点核查填料来源、含水率控制、松铺厚度、碾压遍数、压实度检测频率、雨季施工措施和临时排水措施。"]),
            ("风险处置", ["若方案未说明特殊路基处理、试验段参数、质量检测计划或安全风险控制，应建议退回修改。"]),
        ],
    ),
    DocSpec(
        "DOC-006",
        "02_contract_and_supervision",
        "06_项目资料命名与归档规则.docx",
        "项目资料命名与归档规则",
        "archive_rule",
        "project-context",
        [
            ("命名规则", ["资料名称建议采用 项目-合同段-队伍-资料类型-日期-版本 的格式。"]),
            ("归档要求", ["正式资料应保留签章页、审批页、附件清单和电子文件来源。扫描件应确保文字可识别。"]),
            ("版本控制", ["依据文件、施工方案和审查意见应标识版本状态，过期版本不得作为正式审查依据。"]),
        ],
    ),
    DocSpec(
        "DOC-007",
        "03_team_lj01_opening_condition",
        "07_LJ-01分包合同范围说明.docx",
        "LJ-01 分包合同范围说明",
        "subcontract_scope",
        "opening-condition-lj01",
        [
            ("作业范围", ["LJ-01 作业队承担 K12+000-K18+500 路基清表、挖方、填筑、便道维护和临时排水工程。"]),
            ("接口边界", ["桥梁下部结构、交安设施和永久排水结构不属于 LJ-01 第一阶段作业范围。"]),
            ("监理关注点", ["开工审查应避免将其他队伍的人员、设备或方案误召回到 LJ-01 审查结论中。"]),
        ],
    ),
    DocSpec(
        "DOC-008",
        "03_team_lj01_opening_condition",
        "08_LJ-01企业资质与安全生产许可证摘录.docx",
        "LJ-01 企业资质与安全生产许可证摘录",
        "team_qualification",
        "opening-condition-lj01",
        [
            ("企业信息", ["单位名称：南江恒通土石方工程有限公司。资质类别：施工劳务不分等级。安全生产许可证在模拟有效期内。"]),
            ("适用性说明", ["资质适用于普通路基土石方劳务作业，不覆盖爆破、桥梁专业承包和交通安全设施专业施工。"]),
            ("审查提示", ["若施工方案包含爆破、深基坑或桥梁下构内容，应提示超出 LJ-01 当前资质和合同范围。"]),
        ],
    ),
    DocSpec(
        "DOC-009",
        "03_team_lj01_opening_condition",
        "09_LJ-01项目负责人任命书.docx",
        "LJ-01 项目负责人任命书",
        "personnel",
        "opening-condition-lj01",
        [
            ("任命信息", ["任命李建华为 LJ-01 路基土石方分包作业队现场负责人，负责施工组织、安全管理、质量配合和资料报审。"]),
            ("到岗要求", ["现场负责人应参加监理例会、专项方案交底和开工条件核查。"]),
        ],
    ),
    DocSpec(
        "DOC-014",
        "03_team_lj01_opening_condition",
        "14_LJ-01开工申请表.docx",
        "LJ-01 开工申请表",
        "opening_application",
        "opening-condition-lj01",
        [
            ("申请事项", ["申请 LJ-01 路基土石方分包作业队在 K12+000-K18+500 范围内开展清表、挖方、填筑和临时排水施工。"]),
            ("附件清单", ["已提交资质、安全许可证、人员证书、设备清单、施工方案、安全教育、技术交底、临时排水和便道准备说明。"]),
            ("施工单位自检意见", ["施工单位认为人员、设备、材料、方案和现场准备满足开工条件，申请监理审查。"]),
        ],
    ),
    DocSpec(
        "DOC-015",
        "03_team_lj01_opening_condition",
        "15_LJ-01施工组织与人员进场报审表.docx",
        "LJ-01 施工组织与人员进场报审表",
        "opening_attachment",
        "opening-condition-lj01",
        [
            ("组织架构", ["LJ-01 设置现场负责人、安全员、质检员、测量员、机械班组、运输班组和填筑班组。"]),
            ("进场情况", ["首批进场 38 人，其中管理人员 5 人、机械操作人员 12 人、普通作业人员 21 人。"]),
            ("审查提示", ["人员数量满足试点段初期施工需要，但应核查特种作业证书与设备操作岗位是否匹配。"]),
        ],
    ),
    DocSpec(
        "DOC-016",
        "03_team_lj01_opening_condition",
        "16_LJ-01安全技术交底记录.docx",
        "LJ-01 安全技术交底记录",
        "safety_disclosure",
        "opening-condition-lj01",
        [
            ("交底范围", ["交底内容包括机械作业半径、运输车辆倒车指挥、临边防护、雨季排水、夜间施工照明和个人防护用品。"]),
            ("参加人员", ["现场负责人、安全员、挖掘机司机、压路机司机、运输车辆驾驶员和普通作业人员参加交底。"]),
            ("审查提示", ["交底记录应包含交底人、接受交底人、日期、签字和作业范围。"]),
        ],
    ),
    DocSpec(
        "DOC-018",
        "03_team_lj01_opening_condition",
        "18_LJ-01临时排水与便道准备情况说明.docx",
        "LJ-01 临时排水与便道准备情况说明",
        "site_preparation",
        "opening-condition-lj01",
        [
            ("便道准备", ["K12+000-K18+500 主线左侧设置临时施工便道，已完成压实、错车平台和警示标志布设。"]),
            ("临时排水", ["已设置截水沟、临时边沟和沉淀池，雨季施工前需复核排水出口和下游通畅情况。"]),
            ("审查提示", ["若近期有强降雨预警，应将临时排水复核作为开工前条件。"]),
        ],
    ),
    DocSpec(
        "DOC-019",
        "04_team_lj01_construction_plan",
        "19_K12-K18路基填筑施工方案.docx",
        "K12+000-K18+500 路基填筑施工方案",
        "construction_plan",
        "construction-plan-lj01",
        [
            ("工程范围", ["本方案适用于 K12+000-K18+500 段路基填筑施工，主要包括基底处理、填料运输、分层摊铺、碾压、检测和边坡整修。"]),
            ("施工工艺", ["施工流程为测量放样、清表验槽、试验段施工、填料摊铺、含水率调整、碾压成型、压实度检测、边坡整修。"]),
            ("质量控制", ["填筑前应完成试验段，确定松铺厚度、压实机械组合、碾压遍数和最佳含水率控制范围。"]),
            ("安全措施", ["施工现场设置车辆行驶路线、倒车指挥、机械作业警戒区和夜间照明，雨季施工加强边沟排水。"]),
        ],
        [
            ("主要施工参数", ["控制项", "模拟参数", "审查要点"], [["松铺厚度", "30cm", "需结合试验段确认"], ["碾压遍数", "静压1遍+振压4遍+静压1遍", "需和机械组合匹配"], ["检测频率", "每层每作业段抽检", "需与质量计划一致"]]),
        ],
    ),
    DocSpec(
        "DOC-020",
        "04_team_lj01_construction_plan",
        "20_路基填料来源与试验报告摘要.docx",
        "路基填料来源与试验报告摘要",
        "material_test",
        "construction-plan-lj01",
        [
            ("填料来源", ["填料来自 K15+200 取土场和路堑挖方可利用料，禁止使用淤泥、冻土、有机质土和含草皮树根土。"]),
            ("试验摘要", ["模拟试验显示填料最大干密度、最佳含水率和液塑限指标满足一般路基填筑要求。"]),
            ("审查提示", ["若实际含水率偏离最佳含水率，应采取晾晒、洒水或换填措施。"]),
        ],
    ),
    DocSpec(
        "DOC-021",
        "04_team_lj01_construction_plan",
        "21_路基填筑质量控制与检测计划.docx",
        "路基填筑质量控制与检测计划",
        "quality_plan",
        "construction-plan-lj01",
        [
            ("质量控制点", ["基底处理、填料质量、松铺厚度、含水率、碾压遍数、压实度、宽度、横坡和边坡坡率为主要控制点。"]),
            ("检测安排", ["每层填筑完成后由施工单位自检，监理按规定频率抽检，检测不合格不得进入下一层施工。"]),
            ("整改闭环", ["不合格段应标识范围、分析原因、采取补压或换填措施，并形成复检记录。"]),
        ],
    ),
    DocSpec(
        "DOC-022",
        "04_team_lj01_construction_plan",
        "22_路基施工安全风险辨识与控制措施.docx",
        "路基施工安全风险辨识与控制措施",
        "safety_plan",
        "construction-plan-lj01",
        [
            ("主要风险", ["机械伤害、车辆碰撞、边坡坍塌、雨季冲刷、夜间视线不足和临时用电风险为本阶段主要风险。"]),
            ("控制措施", ["机械作业区设置警戒线，运输车辆限速，雨后复查边坡和便道，夜间施工设置照明和反光标识。"]),
            ("应急要求", ["现场应配备应急联系人、救援路线、排水设备和临时交通疏导措施。"]),
        ],
    ),
    DocSpec(
        "DOC-023",
        "05_review_examples",
        "23_历史样例_开工条件审查通过案例.docx",
        "历史样例：开工条件审查通过案例",
        "review_example",
        "opening-condition-lj01",
        [
            ("案例结论", ["某路基作业队因资质、人员、设备、方案、安全教育和现场准备均齐备，经监理审查同意开工。"]),
            ("通过理由", ["关键岗位到岗，主要机械设备已进场，安全技术交底完整，临时排水和便道具备施工条件。"]),
            ("可复用提示", ["审查通过必须说明依据、资料来源和人工确认结论，不能只依据模型回答。"]),
        ],
    ),
    DocSpec(
        "DOC-024",
        "05_review_examples",
        "24_历史样例_施工方案审查退回修改案例.docx",
        "历史样例：施工方案审查退回修改案例",
        "review_example",
        "construction-plan-lj01",
        [
            ("退回原因", ["某路基填筑方案未说明试验段参数、雨季排水措施和压实度检测频率，被监理退回修改。"]),
            ("整改要求", ["补充试验段成果、机械组合、质量检测计划、雨季施工组织和安全风险控制措施。"]),
            ("可复用提示", ["方案审查应输出缺失项和修改建议，但正式退回意见需由监理工程师确认。"]),
        ],
    ),
    DocSpec(
        "SRC-001",
        "06_source_norms",
        "SRC-001_公路工程施工监理规范_开工与方案审查摘要.docx",
        "公路工程施工监理规范：开工与方案审查摘要",
        "official_source_summary",
        "project-context",
        [
            ("来源", ["标准号：JTG G10-2016。发布单位：交通运输部。官方来源登记见 docs/construction-supervision-pilot-dataset/source-register.md。"]),
            ("审查用途", ["用于提示监理在开工条件、施工方案、质量安全和资料归档方面的审查职责。"]),
            ("入库边界", ["本文件为自写摘要，不替代正式标准文本。正式引用时应回到官方发布文件。"]),
        ],
    ),
    DocSpec(
        "SRC-002",
        "06_source_norms",
        "SRC-002_质量检验评定标准_土建工程评定框架摘要.docx",
        "质量检验评定标准：土建工程评定框架摘要",
        "official_source_summary",
        "construction-plan-lj01",
        [
            ("来源", ["标准号：JTG F80/1-2017。发布单位：交通运输部。"]),
            ("审查用途", ["用于辅助理解单位工程、分部工程、分项工程和检验项目之间的关系。"]),
            ("入库边界", ["本摘要仅帮助召回质量控制点，不能替代正式检验评定标准。"]),
        ],
    ),
    DocSpec(
        "SRC-003",
        "06_source_norms",
        "SRC-003_施工安全技术规范_风险控制摘要.docx",
        "施工安全技术规范：风险控制摘要",
        "official_source_summary",
        "construction-plan-lj01",
        [
            ("来源", ["标准号：JTG F90-2015。发布单位：交通运输部。"]),
            ("审查用途", ["用于辅助识别机械作业、运输、临边、雨季施工和临时用电等安全风险。"]),
            ("入库边界", ["本摘要只用于知识库验证，不作为正式安全专项方案审批依据。"]),
        ],
    ),
    DocSpec(
        "SRC-004",
        "06_source_norms",
        "SRC-004_路基施工技术规范_填筑控制摘要.docx",
        "路基施工技术规范：填筑控制摘要",
        "official_source_summary",
        "construction-plan-lj01",
        [
            ("来源", ["标准号：JTG/T 3610-2019。发布单位：交通运输部。"]),
            ("审查用途", ["用于辅助召回路基填筑的填料、松铺厚度、碾压、压实和检测控制点。"]),
            ("入库边界", ["正式审查时需核对官方规范和项目合同技术条款。"]),
        ],
    ),
    DocSpec(
        "SRC-006",
        "06_source_norms",
        "SRC-006_标准施工招标文件_合同与技术规范摘要.docx",
        "公路工程标准施工招标文件：合同与技术规范摘要",
        "official_source_summary",
        "project-context",
        [
            ("来源", ["文件：公路工程标准施工招标文件 2018 年版。发布单位：交通运输部。"]),
            ("审查用途", ["用于辅助理解合同、技术规范、计量、资料和承包人责任边界。"]),
            ("入库边界", ["项目专用合同条款优先于通用模板摘要。"]),
        ],
    ),
]


SHEETS = [
    SheetSpec(
        "DOC-010",
        "03_team_lj01_opening_condition",
        "10_LJ-01安全员与质检员证书清单.xlsx",
        "LJ-01 安全员与质检员证书清单",
        "personnel",
        "opening-condition-lj01",
        [
            ("人员证书", ["姓名", "岗位", "证书编号", "有效期", "到岗状态", "审查意见"], [["王强", "专职安全员", "AQ-2026-LJ01-001", "2028-06-30", "已到岗", "证书有效"], ["赵敏", "质检员", "ZJ-2026-LJ01-002", "2027-12-31", "已到岗", "证书有效"], ["陈磊", "测量员", "CL-2026-LJ01-003", "2027-09-30", "已到岗", "需核对原件"]]),
        ],
    ),
    SheetSpec(
        "DOC-011",
        "03_team_lj01_opening_condition",
        "11_LJ-01特种作业人员清单.xlsx",
        "LJ-01 特种作业人员清单",
        "personnel",
        "opening-condition-lj01",
        [
            ("特种作业人员", ["姓名", "工种", "证书编号", "设备/作业", "有效期", "审查状态"], [["刘海", "挖掘机司机", "TZ-LJ01-001", "挖掘机", "2028-01-31", "通过"], ["周亮", "压路机司机", "TZ-LJ01-002", "压路机", "2028-03-31", "通过"], ["孙杰", "装载机司机", "TZ-LJ01-003", "装载机", "2027-11-30", "通过"], ["吴刚", "电工", "TZ-LJ01-004", "临时用电", "2027-10-31", "需现场核验"]]),
        ],
    ),
    SheetSpec(
        "DOC-012",
        "03_team_lj01_opening_condition",
        "12_LJ-01主要机械设备进场清单.xlsx",
        "LJ-01 主要机械设备进场清单",
        "equipment",
        "opening-condition-lj01",
        [
            ("设备清单", ["设备名称", "规格型号", "数量", "进场日期", "状态", "用途"], [["挖掘机", "PC220", 3, "2026-07-10", "已进场", "清表与挖方"], ["推土机", "SD16", 2, "2026-07-11", "已进场", "摊铺整平"], ["压路机", "22t", 2, "2026-07-12", "已进场", "路基碾压"], ["洒水车", "12m3", 1, "2026-07-12", "已进场", "含水率调整"], ["自卸车", "20m3", 12, "2026-07-13", "已进场", "填料运输"]]),
        ],
    ),
    SheetSpec(
        "DOC-013",
        "03_team_lj01_opening_condition",
        "13_LJ-01测量与试验仪器校验证明清单.xlsx",
        "LJ-01 测量与试验仪器校验证明清单",
        "calibration",
        "opening-condition-lj01",
        [
            ("仪器清单", ["仪器名称", "编号", "校验单位", "校验日期", "有效期", "用途"], [["全站仪", "TS-LJ01-01", "南江计量检测中心", "2026-06-20", "2027-06-19", "测量放样"], ["水准仪", "DS-LJ01-02", "南江计量检测中心", "2026-06-20", "2027-06-19", "高程控制"], ["灌砂筒", "GS-LJ01-03", "南江计量检测中心", "2026-06-25", "2027-06-24", "压实度检测"], ["电子天平", "TP-LJ01-04", "南江计量检测中心", "2026-06-25", "2027-06-24", "含水率试验"]]),
        ],
    ),
    SheetSpec(
        "DOC-017",
        "03_team_lj01_opening_condition",
        "17_LJ-01三级安全教育记录.xlsx",
        "LJ-01 三级安全教育记录",
        "safety_training",
        "opening-condition-lj01",
        [
            ("安全教育记录", ["姓名", "班组", "公司级", "项目级", "班组级", "考试成绩", "签字状态"], [["刘海", "机械班", "已完成", "已完成", "已完成", 92, "已签字"], ["周亮", "机械班", "已完成", "已完成", "已完成", 90, "已签字"], ["孙杰", "运输班", "已完成", "已完成", "已完成", 88, "已签字"], ["吴刚", "电工", "已完成", "已完成", "已完成", 91, "已签字"], ["李明", "填筑班", "已完成", "已完成", "已完成", 86, "已签字"]]),
        ],
    ),
]


def write_manifest(entries: list[dict[str, str]]) -> None:
    manifest = {
        "dataset_id": "dataset-njdl-jd-a1-pilot",
        "dataset_name": "南江至东岭高速公路改扩建工程 JD-A1 监理审查模拟资料包",
        "grain": "project-level-logical-knowledge-base-with-metadata-filtering",
        "generated_by": "docs/construction-supervision-pilot-dataset/scripts/generate_pilot_dataset.py",
        "metadata": BASE_META,
        "documents": entries,
    }
    manifest_path = ROOT / "00_manifest" / "dataset-manifest.json"
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")

    yaml_lines = [
        "dataset_id: dataset-njdl-jd-a1-pilot",
        "dataset_name: 南江至东岭高速公路改扩建工程 JD-A1 监理审查模拟资料包",
        "grain: project-level-logical-knowledge-base-with-metadata-filtering",
        "generated_by: docs/construction-supervision-pilot-dataset/scripts/generate_pilot_dataset.py",
        "metadata:",
    ]
    for key, value in BASE_META.items():
        yaml_lines.append(f"  {key}: {value}")
    yaml_lines.append("documents:")
    for entry in entries:
        yaml_lines.append(f"  - document_id: {entry['document_id']}")
        yaml_lines.append(f"    relative_path: {entry['relative_path']}")
        yaml_lines.append(f"    file_format: {entry['file_format']}")
        yaml_lines.append(f"    document_type: {entry['document_type']}")
        yaml_lines.append(f"    review_task_id: {entry['review_task_id']}")
        yaml_lines.append(f"    content_hash: {entry['content_hash']}")
    (ROOT / "00_manifest" / "dataset-manifest.yaml").write_text("\n".join(yaml_lines) + "\n", encoding="utf-8")

    csv_path = ROOT / "00_manifest" / "dataset-manifest.csv"
    headers = [
        "document_id",
        "relative_path",
        "file_format",
        "document_type",
        "review_task_id",
        "project_id",
        "contract_package_id",
        "subcontract_team_id",
        "basis_version_id",
        "content_hash",
    ]
    lines = [",".join(headers)]
    for entry in entries:
        lines.append(",".join(str(entry.get(header, "")).replace(",", "，") for header in headers))
    csv_path.write_text("\n".join(lines) + "\n", encoding="utf-8")

    wb = Workbook()
    ws = wb.active
    ws.title = "manifest"
    ws.append(headers)
    for entry in entries:
        ws.append([entry.get(header, "") for header in headers])
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E79")
    ws.freeze_panes = "A2"
    autosize(ws)
    wb.save(ROOT / "00_manifest" / "dataset-manifest.xlsx")


def main() -> None:
    ensure_dirs()
    entries = []
    for doc_spec in DOCS:
        entries.append(write_docx(doc_spec))
    for sheet_spec in SHEETS:
        entries.append(write_xlsx(sheet_spec))
    write_manifest(entries)
    print(f"Generated {len(entries)} pilot documents under {ROOT}")


if __name__ == "__main__":
    main()
