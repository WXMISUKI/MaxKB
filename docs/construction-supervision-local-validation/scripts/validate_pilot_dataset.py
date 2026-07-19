# coding=utf-8
"""
Validate the simulated construction-supervision pilot dataset and generate MaxKB upload aids.

This script is intentionally independent from MaxKB internals. It validates the local pilot data pack before manual
upload through the MaxKB UI.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
from collections import Counter
from pathlib import Path
from typing import Any

from docx import Document
from openpyxl import load_workbook


REQUIRED_FIELDS = [
    "organization_id",
    "project_id",
    "contract_package_id",
    "section_id",
    "supervision_section_id",
    "subcontract_team_id",
    "review_task_id",
    "basis_version_id",
    "document_type",
    "source_object_id",
    "content_hash",
    "relative_path",
    "file_format",
]

UPLOAD_GROUPS = {
    "01_project_basis": "01 项目依据",
    "02_contract_and_supervision": "02 合同与监理依据",
    "03_team_lj01_opening_condition": "03 LJ-01 开工条件资料",
    "04_team_lj01_construction_plan": "04 LJ-01 施工方案资料",
    "05_review_examples": "05 历史审查样例",
    "06_source_norms": "06 官方规范摘要",
}

OPENING_QUESTIONS = [
    "请基于知识库资料，列出 LJ-01 路基土石方分包作业队开工条件审查需要核对的材料清单，并标注引用来源。",
    "LJ-01 队伍的企业资质和安全生产许可证是否覆盖 K12+000-K18+500 路基土石方作业？有哪些不应扩展到本队伍的作业范围？",
    "LJ-01 开工条件审查中，项目负责人、安全员、质检员、测量员是否已经到岗？需要人工复核哪些证书原件？",
    "请核查 LJ-01 特种作业人员与主要机械设备是否匹配，指出需要现场核验的项目。",
    "LJ-01 的主要机械设备、测量试验仪器、安全教育、技术交底和临时排水便道准备是否支持开工？请按缺失项/风险项输出。",
    "请给出 LJ-01 开工条件审查的支持性召回摘要。注意不要直接作出正式批准结论，最终结论需人工确认。",
]

PLAN_QUESTIONS = [
    "请判断 K12+000-K18+500 路基填筑施工方案是否覆盖 LJ-01 的合同作业范围，并列出依据来源。",
    "路基填筑施工方案是否说明施工流程、试验段、松铺厚度、碾压遍数、含水率控制和压实度检测？缺失项有哪些？",
    "请结合填料来源与试验报告摘要，判断填料适用性审查需要关注哪些指标和风险。",
    "路基填筑质量控制与检测计划是否能支持分层填筑、压实度检测和不合格整改闭环？",
    "路基施工安全风险辨识是否覆盖机械伤害、车辆运输、临边、雨季排水、夜间施工和临时用电？",
    "请输出施工方案审查的退回修改建议模板，要求每条建议附知识库引用来源，并保留人工复核入口。",
]


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as file:
        for chunk in iter(lambda: file.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def read_manifest(root: Path) -> dict[str, Any]:
    manifest_path = root / "00_manifest" / "dataset-manifest.json"
    if not manifest_path.exists():
        raise FileNotFoundError(f"Manifest not found: {manifest_path}")
    return json.loads(manifest_path.read_text(encoding="utf-8-sig"))


def validate_docx(path: Path) -> tuple[bool, str]:
    try:
        doc = Document(path)
        paragraphs = [paragraph.text.strip() for paragraph in doc.paragraphs if paragraph.text.strip()]
        if not paragraphs:
            return False, "docx has no non-empty paragraphs"
        return True, f"paragraphs={len(paragraphs)}"
    except Exception as exc:  # pragma: no cover - defensive CLI boundary
        return False, str(exc)


def validate_xlsx(path: Path) -> tuple[bool, str]:
    try:
        workbook = load_workbook(path, read_only=True, data_only=False)
        sheet_count = len(workbook.sheetnames)
        non_empty = 0
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                if any(cell.value not in (None, "") for cell in row):
                    non_empty += 1
                    break
        workbook.close()
        if non_empty == 0:
            return False, "xlsx has no non-empty sheets"
        return True, f"sheets={sheet_count}, non_empty_sheets={non_empty}"
    except Exception as exc:  # pragma: no cover - defensive CLI boundary
        return False, str(exc)


def upload_group(relative_path: str) -> str:
    folder = relative_path.split("/", 1)[0]
    return UPLOAD_GROUPS.get(folder, "99 未分类")


def validate_dataset(root: Path) -> dict[str, Any]:
    manifest = read_manifest(root)
    documents = manifest.get("documents", [])
    errors: list[str] = []
    warnings: list[str] = []
    rows: list[dict[str, Any]] = []
    counts = Counter()
    by_task = Counter()
    by_type = Counter()

    seen_paths: set[str] = set()
    for index, entry in enumerate(documents, start=1):
        doc_id = entry.get("document_id", f"row-{index}")
        missing = [field for field in REQUIRED_FIELDS if not entry.get(field)]
        if missing:
            errors.append(f"{doc_id}: missing required fields: {', '.join(missing)}")

        relative_path = entry.get("relative_path", "")
        if relative_path in seen_paths:
            errors.append(f"{doc_id}: duplicated relative_path: {relative_path}")
        seen_paths.add(relative_path)

        path = root / relative_path
        if not path.exists():
            errors.append(f"{doc_id}: file not found: {relative_path}")
            continue

        actual_hash = sha256(path)
        if entry.get("content_hash") != actual_hash:
            errors.append(f"{doc_id}: content_hash mismatch for {relative_path}")

        suffix = path.suffix.lower().lstrip(".")
        if suffix != entry.get("file_format"):
            errors.append(f"{doc_id}: file_format mismatch, manifest={entry.get('file_format')} actual={suffix}")

        if suffix == "docx":
            ok, detail = validate_docx(path)
        elif suffix == "xlsx":
            ok, detail = validate_xlsx(path)
        else:
            ok, detail = False, f"unsupported upload format: {suffix}"
        if not ok:
            errors.append(f"{doc_id}: unreadable {suffix}: {detail}")

        group = upload_group(relative_path)
        counts[suffix] += 1
        by_task[entry.get("review_task_id", "")] += 1
        by_type[entry.get("document_type", "")] += 1
        rows.append(
            {
                "upload_order": index,
                "upload_group": group,
                "document_id": doc_id,
                "relative_path": relative_path,
                "file_format": suffix,
                "document_type": entry.get("document_type", ""),
                "review_task_id": entry.get("review_task_id", ""),
                "validation_detail": detail,
            }
        )

    all_files = {
        str(path.relative_to(root)).replace("\\", "/")
        for path in root.rglob("*")
        if path.is_file() and path.suffix.lower() in {".docx", ".xlsx"} and "00_manifest" not in path.parts
    }
    manifest_files = {entry.get("relative_path", "") for entry in documents}
    unmanaged = sorted(all_files - manifest_files)
    if unmanaged:
        warnings.append(f"{len(unmanaged)} Office files are not listed as upload documents: {', '.join(unmanaged)}")

    for required_task in ["opening-condition-lj01", "construction-plan-lj01"]:
        if by_task[required_task] == 0:
            errors.append(f"missing review_task_id coverage: {required_task}")

    return {
        "dataset_id": manifest.get("dataset_id"),
        "dataset_name": manifest.get("dataset_name"),
        "root": str(root),
        "document_count": len(documents),
        "format_counts": dict(counts),
        "review_task_counts": dict(by_task),
        "document_type_counts": dict(by_type),
        "errors": errors,
        "warnings": warnings,
        "upload_rows": rows,
    }


def write_upload_plan(root: Path, rows: list[dict[str, Any]]) -> None:
    target = root / "00_manifest" / "maxkb-upload-plan.csv"
    headers = [
        "upload_order",
        "upload_group",
        "document_id",
        "relative_path",
        "file_format",
        "document_type",
        "review_task_id",
    ]
    with target.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=headers)
        writer.writeheader()
        for row in rows:
            writer.writerow({header: row.get(header, "") for header in headers})


def write_question_set(root: Path) -> None:
    target = root / "00_manifest" / "review-question-set.md"
    lines = [
        "# MaxKB 本地验收问题集",
        "",
        "## 开工条件审查",
        "",
    ]
    for index, question in enumerate(OPENING_QUESTIONS, start=1):
        lines.append(f"{index}. {question}")
    lines.extend(["", "## 施工方案审查", ""])
    for index, question in enumerate(PLAN_QUESTIONS, start=1):
        lines.append(f"{index}. {question}")
    lines.extend(
        [
            "",
            "## 记录要求",
            "",
            "- 每个问题记录 Top 召回来源、遗漏来源、误召回来源和人工判断。",
            "- 输出不得直接作为正式批准或退回结论。",
            "- 若无法按项目、队伍、审查任务过滤，需要进入下一轮 metadata 扩展设计。",
        ]
    )
    target.write_text("\n".join(lines) + "\n", encoding="utf-8")


def write_report(root: Path, result: dict[str, Any]) -> None:
    report_json = {key: value for key, value in result.items() if key != "upload_rows"}
    (root / "00_manifest" / "validation-report.json").write_text(
        json.dumps(report_json, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    lines = [
        "# 模拟资料包校验报告",
        "",
        f"- 数据集：{result.get('dataset_name')}",
        f"- 文档数量：{result.get('document_count')}",
        f"- 格式统计：{result.get('format_counts')}",
        f"- 审查任务覆盖：{result.get('review_task_counts')}",
        f"- 阻塞错误：{len(result.get('errors', []))}",
        f"- 警告：{len(result.get('warnings', []))}",
        "",
        "## 阻塞错误",
        "",
    ]
    errors = result.get("errors", [])
    if errors:
        lines.extend(f"- {error}" for error in errors)
    else:
        lines.append("- 无")
    lines.extend(["", "## 警告", ""])
    warnings = result.get("warnings", [])
    if warnings:
        lines.extend(f"- {warning}" for warning in warnings)
    else:
        lines.append("- 无")
    lines.extend(
        [
            "",
            "## 下一步",
            "",
            "1. 在 MaxKB 中创建项目级知识库。",
            "2. 按 `maxkb-upload-plan.csv` 上传资料。",
            "3. 使用 `review-question-set.md` 执行开工条件和施工方案审查问题。",
            "4. 记录召回质量，再决定是否进入源码级 metadata 扩展或工作流原型。",
        ]
    )
    (root / "00_manifest" / "validation-report.md").write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> int:
    parser = argparse.ArgumentParser(description="Validate construction supervision pilot dataset.")
    parser.add_argument(
        "--root",
        default="docs/simulated-pilot-dataset/NJDL-JD-A1",
        help="Pilot dataset root containing 00_manifest/dataset-manifest.json",
    )
    args = parser.parse_args()

    root = Path(args.root).resolve()
    result = validate_dataset(root)
    write_upload_plan(root, result["upload_rows"])
    write_question_set(root)
    write_report(root, result)

    print(f"dataset={result.get('dataset_name')}")
    print(f"documents={result.get('document_count')}")
    print(f"errors={len(result.get('errors', []))}")
    print(f"warnings={len(result.get('warnings', []))}")
    return 1 if result.get("errors") else 0


if __name__ == "__main__":
    raise SystemExit(main())
